from ..model import Assumptions, OpeningBalance

def from_excel(cls, path):
    """Import the toolkit's versioned workbook; arbitrary Excel models are unsupported."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        def pairs(sheet):
            result = {}
            for row in wb[sheet].iter_rows(values_only=True):
                key, value = row[:2]
                if key is None:
                    continue
                if key in result:
                    raise ValueError(f"Duplicate workbook input: {key}")
                if isinstance(value, str) and value.startswith("="):
                    raise ValueError("Input formulas are unsupported; supply numeric values")
                result[key] = value
            return result
        inputs = pairs("Assumptions")
        if inputs.pop("finmodel_schema") != 1:
            raise ValueError("Unsupported workbook schema")
        name, years = inputs.pop("name"), inputs.pop("years")
        drivers = {}
        if "Drivers" in wb.sheetnames:
            for year, key, value in wb["Drivers"].iter_rows(min_row=2, values_only=True):
                if year is None and key is None and value is None:
                    continue
                if isinstance(value, str) and value.startswith("="):
                    raise ValueError("Input formulas are unsupported")
                if key in drivers.setdefault(year, {}):
                    raise ValueError("Duplicate driver input")
                drivers[year][key] = value
        return cls(assumptions=Assumptions(**inputs), opening=OpeningBalance(**pairs("Opening")), name=name).forecast(years, drivers=drivers)
    finally:
        wb.close()
