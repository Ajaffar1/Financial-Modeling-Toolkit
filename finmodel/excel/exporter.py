from dataclasses import asdict

def export_excel(model, path):
    from openpyxl import Workbook
    if not model.periods:
        raise ValueError("Call forecast() first")
    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    ws.append(["finmodel_schema", 1])
    ws.append(["name", model.name])
    ws.append(["years", len(model.periods)])
    for k, v in asdict(model.assumptions).items():
        ws.append([k, v])
    ws = wb.create_sheet("Drivers")
    ws.append(["Year", "Assumption", "Value"])
    for year, overrides in sorted(model.drivers.items()):
        for key, value in overrides.items():
            ws.append([year, key, value])
    ws = wb.create_sheet("Opening")
    for k, v in asdict(model.opening).items():
        ws.append([k, v])
    for label, attribute in [("Income Statement", "income_statement"), ("Balance Sheet", "balance_sheet"), ("Cash Flow", "cash_flow")]:
        ws = wb.create_sheet(label)
        fields = list(asdict(getattr(model.periods[0], attribute)))
        ws.append(["Year", *fields])
        for p in model.periods:
            ws.append([p.year, *asdict(getattr(p, attribute)).values()])
    ws = wb.create_sheet("Valuation")
    ws.append(["Enterprise value", model.enterprise_value()])
    ws.append(["Equity value", model.equity_value()])
    ws = wb.create_sheet("Audit")
    ws.append(["Code", "Severity", "Message", "Year"])
    for finding in model.audit():
        ws.append(list(asdict(finding).values()))
    for ws in wb:
        ws.freeze_panes = "B2"
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(65, max(18, max(len(str(c.value or "")) for c in column) + 2))
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
    wb.save(path)
    return path
